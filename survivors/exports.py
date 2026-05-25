"""تصدير واستيراد بيانات الناجين: Excel, PDF (HTML print), JSON."""

import io
import json
from datetime import date, datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.serializers import deserialize, serialize
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.translation import gettext_lazy as _

from accounts.models import AuditLog

from .forms import JSONImportForm, SurvivorSearchForm
from .models import (
    ChainOfCustodyLog, DetentionEvent, DetentionFacility, DetentionPeriod,
    InformedConsent, Interview, InterviewMedia, LongTermImpact,
    MedicalAssessment, ReleaseEvent, SupportingDocument, SurvivorNote,
    SurvivorProfile, TortureMethod, Witness,
)


# ============================================================
# الأعمدة القابلة للتصدير (Excel/PDF)
# ============================================================

AVAILABLE_COLUMNS = [
    # (key, label, getter)
    ("case_reference", _("رقم القضية"), lambda s: s.case_reference),
    ("full_name", _("الاسم الكامل"), lambda s: s.full_name),
    ("first_name", _("الاسم الأول"), lambda s: s.first_name),
    ("father_name", _("اسم الأب"), lambda s: s.father_name),
    ("family_name", _("اسم العائلة"), lambda s: s.family_name),
    ("mother_name", _("اسم الأم"), lambda s: s.mother_name),
    ("alias", _("الكنية"), lambda s: s.alias),
    ("national_id", _("الرقم الوطني"), lambda s: s.national_id),
    ("birth_date", _("تاريخ الميلاد"), lambda s: s.birth_date),
    ("birth_governorate", _("محافظة الولادة"), lambda s: s.get_birth_governorate_display()),
    ("gender", _("الجنس"), lambda s: s.get_gender_display()),
    ("marital_status", _("الحالة الزوجية"), lambda s: s.get_marital_status_at_detention_display()),
    ("occupation", _("المهنة"), lambda s: s.get_occupation_category_display()),
    ("occupation_detail", _("تفاصيل المهنة"), lambda s: s.occupation_detail),
    ("political_activity", _("النشاط السياسي"), lambda s: s.get_political_activity_category_display()),
    ("governorate_at_detention", _("محافظة الاعتقال"), lambda s: s.get_governorate_at_detention_display()),
    ("current_country", _("بلد الإقامة"), lambda s: s.get_current_country_display()),
    ("current_governorate", _("المحافظة الحالية"), lambda s: s.get_current_governorate_display()),
    ("current_city", _("المدينة الحالية"), lambda s: s.current_city),
    ("current_phone", _("الهاتف"), lambda s: s.current_phone),
    ("current_email", _("البريد"), lambda s: s.current_email),
    ("age_at_detention", _("العمر وقت الاعتقال"), lambda s: s.age_at_detention),
    ("first_detention_date", _("تاريخ أول اعتقال"),
     lambda s: (s.detention_events.order_by("detention_date").first() or _O()).detention_date),
    ("release_date", _("تاريخ الإفراج"),
     lambda s: getattr(getattr(s, "release_event", None), "release_date", None)),
    ("release_type", _("نوع الإفراج"),
     lambda s: (getattr(s, "release_event", None) and s.release_event.get_release_type_display())),
    ("total_detention_periods", _("عدد فترات الاحتجاز"),
     lambda s: s.detention_periods.count()),
    ("total_detention_days", _("إجمالي أيام الاحتجاز"),
     lambda s: sum((p.duration_days or 0) for p in s.detention_periods.all() if p.duration_days)),
    ("first_facility", _("أول فرع"),
     lambda s: str((s.detention_periods.order_by("order_index", "from_date").first() or _O()).facility or "")),
    ("classification", _("التصنيف"), lambda s: s.get_file_classification_display()),
    ("reliability_score", _("الموثوقية"), lambda s: s.reliability_score),
    ("corroboration_score", _("التحقق المتقاطع"), lambda s: s.corroboration_score),
    ("completeness_score", _("الاكتمال"), lambda s: s.completeness_score),
    ("overall_score", _("الدرجة الإجمالية"), lambda s: s.overall_score),
    ("witnesses_count", _("عدد الشهود"), lambda s: s.witnesses.count()),
    ("independent_witnesses", _("شهود مستقلون"), lambda s: s.independent_witnesses_count),
    ("documents_count", _("عدد الوثائق"), lambda s: s.documents.count()),
    ("interviews_count", _("عدد المقابلات"), lambda s: s.interviews.count()),
    ("videos_count", _("عدد الفيديوهات"), lambda s: s.total_videos),
    ("has_consent", _("لديه موافقة"),
     lambda s: "نعم" if (getattr(s, "consent", None) and s.consent.is_fully_compliant) else "لا"),
    ("consent_iiim", _("موافقة IIIM"),
     lambda s: "نعم" if (getattr(s, "consent", None) and s.consent.share_with_iiim) else "لا"),
    ("consent_icc", _("موافقة ICC"),
     lambda s: "نعم" if (getattr(s, "consent", None) and s.consent.share_with_icc) else "لا"),
    ("has_medical", _("تقييم طبي"),
     lambda s: "نعم" if s.medical_assessments.exists() else "لا"),
    ("istanbul_compliant", _("متوافق إسطنبول"),
     lambda s: "نعم" if s.medical_assessments.filter(istanbul_protocol_compliant=True).exists() else "لا"),
    # حقول المسح الاجتماعي
    ("household_size", _("حجم الأسرة"),
     lambda s: getattr(getattr(s, "household_survey", None), "household_size", None)),
    ("children_count", _("عدد الأبناء"),
     lambda s: getattr(getattr(s, "household_survey", None), "children_count", None)),
    ("marital_now", _("الحالة الزوجية الحالية"),
     lambda s: (getattr(s, "household_survey", None) and s.household_survey.get_marital_status_display())),
    ("displacement_status", _("التهجير"),
     lambda s: (getattr(s, "household_survey", None) and s.household_survey.get_displacement_status_display())),
    ("housing_type", _("نوع السكن"),
     lambda s: (getattr(getattr(s, "household_survey", None), "housing", None) and s.household_survey.housing.get_housing_type_display())),
    ("rent_amount", _("الإيجار"),
     lambda s: (getattr(getattr(s, "household_survey", None), "housing", None) and s.household_survey.housing.rent_amount)),
    ("employment_status", _("وضع العمل"),
     lambda s: (getattr(s, "employment", None) and s.employment.get_status_display())),
    ("monthly_income", _("الدخل الشهري"),
     lambda s: (getattr(s, "employment", None) and s.employment.monthly_income)),
    ("documenter", _("الموثّق"),
     lambda s: s.documenter.username if s.documenter else ""),
    ("created_at", _("تاريخ الإضافة"), lambda s: s.created_at),
    ("notes_count", _("عدد الملاحظات"), lambda s: s.notes.count()),
]


class _O:
    """عنصر فارغ نُرجع منه None لتجنّب الأخطاء عند None.attribute."""

    def __getattr__(self, name):
        return None


def _column_dict():
    return {key: (label, getter) for key, label, getter in AVAILABLE_COLUMNS}


# ============================================================
# واجهة اختيار الأعمدة + التصدير
# ============================================================

def _filtered_queryset(request):
    """يطبّق الفلاتر من querystring على نفس منطق قائمة الناجين."""
    from .views import _apply_survivor_filters
    form = SurvivorSearchForm(request.GET or None)
    qs = SurvivorProfile.objects.select_related("documenter")
    if form.is_valid():
        qs = _apply_survivor_filters(qs, form.cleaned_data)
    return qs.order_by("case_reference")


DEFAULT_COLUMNS = [
    "case_reference", "full_name", "gender",
    "governorate_at_detention", "classification",
    "overall_score", "first_detention_date", "release_date",
]


@login_required
def export_chooser(request):
    """واجهة لاختيار الأعمدة والصيغة قبل التصدير."""
    qs = _filtered_queryset(request)
    count = qs.count()
    return render(request, "survivors/export_chooser.html", {
        "columns": AVAILABLE_COLUMNS,
        "default_columns": DEFAULT_COLUMNS,
        "default_columns_json": json.dumps(DEFAULT_COLUMNS),
        "count": count,
        "querystring": request.GET.urlencode(),
    })


def _safe(value):
    """يحوّل القيمة لنوع آمن للكتابة في Excel/JSON."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "__str__"):
        return str(value)
    return value


@login_required
def export_excel(request):
    """تصدير Excel بالأعمدة المختارة."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols_param = request.GET.get("columns", "")
    selected_keys = [k for k in cols_param.split(",") if k]
    if not selected_keys:
        selected_keys = [
            "case_reference", "full_name", "gender",
            "governorate_at_detention", "classification",
            "overall_score", "first_detention_date",
        ]

    qs = _filtered_queryset(request)
    col_map = _column_dict()

    wb = Workbook()
    ws = wb.active
    ws.title = "ملفات الناجين"
    ws.sheet_view.rightToLeft = True

    headers = []
    for key in selected_keys:
        if key in col_map:
            headers.append(str(col_map[key][0]))
    ws.append(headers)
    # تنسيق رأس الجدول
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1C6B85")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    # البيانات
    for survivor in qs:
        row = []
        for key in selected_keys:
            if key in col_map:
                getter = col_map[key][1]
                try:
                    val = getter(survivor)
                except Exception:
                    val = ""
                row.append(_safe(val))
        ws.append(row)

    # سجل التدقيق
    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="SurvivorProfile",
        target_repr=f"تصدير Excel ({qs.count()} ملف، {len(headers)} عمود)",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"haqquna_survivors_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def export_pdf_print(request):
    """صفحة HTML جاهزة للطباعة كـPDF عبر المتصفح (دعم RTL/عربي ممتاز)."""
    cols_param = request.GET.get("columns", "")
    selected_keys = [k for k in cols_param.split(",") if k]
    if not selected_keys:
        selected_keys = [
            "case_reference", "full_name", "gender",
            "governorate_at_detention", "classification", "overall_score",
        ]
    qs = _filtered_queryset(request)
    col_map = _column_dict()

    rows = []
    for survivor in qs:
        row = []
        for key in selected_keys:
            if key in col_map:
                try:
                    val = col_map[key][1](survivor)
                except Exception:
                    val = ""
                row.append(_safe(val))
        rows.append(row)

    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="SurvivorProfile",
        target_repr=f"تصدير PDF ({len(rows)} ملف)",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )

    headers = [str(col_map[k][0]) for k in selected_keys if k in col_map]
    return render(request, "survivors/export_print.html", {
        "headers": headers, "rows": rows,
        "count": len(rows), "today": datetime.now(),
    })


# ============================================================
# تصدير JSON كامل للقاعدة
# ============================================================

JSON_MODELS = [
    DetentionFacility, TortureMethod, SurvivorProfile, InformedConsent,
    DetentionEvent, DetentionPeriod, ReleaseEvent, Witness,
    SupportingDocument, ChainOfCustodyLog, MedicalAssessment,
    LongTermImpact, SurvivorNote, Interview, InterviewMedia,
]


@login_required
def export_json(request):
    """تصدير قاعدة بيانات الناجين كاملةً بصيغة JSON."""
    if not (request.user.is_superuser or request.user.role in ("admin", "protection")):
        messages.error(request, _("لا تملك صلاحية تصدير قاعدة البيانات."))
        return redirect("survivors:list")

    all_objects = []
    for model in JSON_MODELS:
        all_objects.extend(list(model.objects.all()))

    data = serialize("json", all_objects, indent=2, ensure_ascii=False)

    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="Database",
        target_repr=f"تصدير JSON كامل ({len(all_objects)} سجل)",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )

    filename = f"haqquna_backup_{datetime.now():%Y%m%d_%H%M%S}.json"
    response = HttpResponse(data, content_type="application/json; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def import_json(request):
    """استيراد بيانات الناجين من ملف JSON مُصدّر مسبقاً."""
    if not (request.user.is_superuser or request.user.role == "admin"):
        messages.error(request, _("لا تملك صلاحية استيراد قاعدة البيانات."))
        return redirect("survivors:list")

    if request.method == "POST":
        form = JSONImportForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded = form.cleaned_data["file"]
            strategy = form.cleaned_data["merge_strategy"]
            try:
                uploaded.seek(0)
                content = uploaded.read().decode("utf-8")
                created = 0
                skipped = 0
                updated = 0
                update_log = []  # لتسجيل ما تم تعديله
                with transaction.atomic():
                    for deserialized in deserialize("json", content):
                        obj = deserialized.object
                        model = type(obj)
                        # حماية: لا نسمح أبداً باستيراد AuditLog أو ChainOfCustody
                        model_label = f"{model._meta.app_label}.{model.__name__}"
                        if model_label in ("accounts.AuditLog",
                                           "survivors.ChainOfCustodyLog"):
                            skipped += 1
                            continue
                        # تحقّق من الوجود
                        if obj.pk and model.objects.filter(pk=obj.pk).exists():
                            if strategy == "skip_existing":
                                skipped += 1
                                continue
                            # نسجل التحديث في AuditLog (قبل الاستبدال)
                            existing = model.objects.get(pk=obj.pk)
                            update_log.append(
                                f"{model_label} pk={obj.pk}: استُبدل '{str(existing)[:80]}'"
                            )
                            deserialized.save()
                            updated += 1
                        else:
                            deserialized.save()
                            created += 1

                AuditLog.objects.create(
                    user=request.user, action=AuditLog.Action.CREATE,
                    target_model="JSONImport",
                    target_repr=f"استيراد: {created} جديد، {updated} محدّث، {skipped} متخطّى",
                    path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
                    notes=("\n".join(update_log))[:5000] if update_log else "",
                )
                messages.success(
                    request,
                    _("تم الاستيراد بنجاح: %(c)d سجل جديد، %(u)d محدّث، %(s)d متخطّى.")
                    % {"c": created, "u": updated, "s": skipped},
                )
                return redirect("survivors:list")
            except Exception as e:
                messages.error(request, _("فشل الاستيراد: %(err)s") % {"err": str(e)})
    else:
        form = JSONImportForm()

    return render(request, "survivors/import_json.html", {"form": form})


# ============================================================
# تصدير ملف ناجٍ مفرد (PDF print)
# ============================================================

@login_required
def survivor_print(request, pk):
    """صفحة جاهزة للطباعة لملف ناجٍ كامل."""
    survivor = get_object_or_404(SurvivorProfile, pk=pk)
    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="SurvivorProfile", target_id=str(survivor.pk),
        target_repr=f"طباعة ملف {survivor.case_reference}",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )
    return render(request, "survivors/survivor_print.html", {
        "survivor": survivor,
        "consent": getattr(survivor, "consent", None),
        "release": getattr(survivor, "release_event", None),
        "impact": getattr(survivor, "long_term_impact", None),
        "detention_events": survivor.detention_events.all(),
        "detention_periods": survivor.detention_periods.select_related("facility").all(),
        "witnesses": survivor.witnesses.select_related("facility_witnessed_at").all(),
        "documents": survivor.documents.all(),
        "medical_assessments": survivor.medical_assessments.all(),
        "notes": survivor.notes.all(),
        "interviews": survivor.interviews.prefetch_related("media").all(),
        "today": datetime.now(),
    })


@login_required
def consent_print(request, pk):
    """نموذج الموافقة المستنيرة جاهز للطباعة وتوقيع الناجي بخط يده."""
    survivor = get_object_or_404(SurvivorProfile, pk=pk)
    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="InformedConsent",
        target_repr=f"طباعة نموذج موافقة {survivor.case_reference}",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )
    return render(request, "survivors/consent_print.html", {
        "survivor": survivor,
        "consent": getattr(survivor, "consent", None),
        "today": datetime.now(),
    })


# ============================================================
# الإحالة لجهات خارجية (ICC, IIIM, محاكم الولاية القضائية العالمية...)
# ============================================================

from .recipients import RECIPIENTS, all_recipients_by_category, get_recipient


def _check_consent_for_recipient(consent, recipient):
    """يتحقق هل الموافقة الموقّعة تسمح بالمشاركة مع المستلم."""
    req = recipient.get("requires_consent")
    if not req:
        return True, None
    if not consent:
        return False, "لا توجد موافقة مستنيرة موثّقة."
    if not consent.consent_documented or consent.consent_withdrawn:
        return False, "الموافقة المستنيرة غير مكتملة أو مسحوبة."
    if not getattr(consent, req, False):
        return False, f"الناجي لم يوافق صراحة على المشاركة مع هذه الجهة (حقل: {req})."
    return True, None


@login_required
def transmit_chooser(request, pk):
    """صفحة اختيار الجهة الخارجية وكتابة ملاحظات الإحالة."""
    survivor = get_object_or_404(SurvivorProfile, pk=pk)
    consent = getattr(survivor, "consent", None)

    categorized = all_recipients_by_category()
    # نُمرّر لكل مستلم حالة الموافقة (متاح/غير متاح)
    for cat, recipients in categorized.items():
        for r in recipients:
            allowed, reason = _check_consent_for_recipient(consent, r)
            r["allowed"] = allowed
            r["reason"] = reason

    return render(request, "survivors/transmit_chooser.html", {
        "survivor": survivor, "consent": consent,
        "categorized": categorized,
    })


def _sanitize_user_text(text, max_len=500):
    """قص + إزالة الأحرف التحكمية - للحماية من حقن النصوص."""
    if not text:
        return ""
    # إزالة أحرف التحكم (ما عدا newline و tab)
    cleaned = "".join(
        ch for ch in str(text)
        if ch == "\n" or ch == "\t" or (ord(ch) >= 32 and ord(ch) != 127)
    )
    return cleaned[:max_len].strip()


@login_required
def transmit_package(request, pk):
    """يولّد حزمة الإحالة: خطاب رسمي + ملف الناجي - جاهز للطباعة كـPDF.

    يستخدم transaction.atomic + select_for_update لإحكام التحقق من الموافقة
    وتجنب سباق سحب الموافقة بين الفحص والإصدار.
    """
    from django.db import transaction

    survivor = get_object_or_404(SurvivorProfile, pk=pk)
    recipient_code = request.GET.get("recipient", "")[:30]
    recipient = get_recipient(recipient_code)
    if not recipient:
        messages.error(request, "الجهة المستلمة غير معروفة.")
        return redirect("survivors:transmit_chooser", pk=pk)

    # نتحقق من الموافقة داخل transaction مع قفل سجل الموافقة
    with transaction.atomic():
        consent = (
            InformedConsent.objects.select_for_update()
            .filter(survivor=survivor).first()
        )
        allowed, reason = _check_consent_for_recipient(consent, recipient)
        if not allowed:
            messages.error(request, f"لا يمكن الإحالة: {reason}")
            return redirect("survivors:transmit_chooser", pk=pk)

        # ملاحظات اختيارية - مُطهّرة
        transmission_note = _sanitize_user_text(request.GET.get("note", ""))
        reference_number = (
            f"HQ-TX-{survivor.case_reference}-{recipient_code.upper()}-{datetime.now():%Y%m%d}"
        )

        AuditLog.objects.create(
            user=request.user, action=AuditLog.Action.EXPORT,
            target_model="SurvivorProfile", target_id=str(survivor.pk),
            target_repr=f"إحالة {survivor.case_reference} → {recipient['name_en']}",
            path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
            notes=f"reference={reference_number} | recipient={recipient_code} | note={transmission_note[:200]}",
        )

    return render(request, "survivors/transmit_package.html", {
        "survivor": survivor,
        "consent": consent,
        "release": getattr(survivor, "release_event", None),
        "recipient": recipient,
        "transmission_note": transmission_note,
        "reference_number": reference_number,
        "documenter": survivor.documenter,
        "transmitted_by": request.user,
        "detention_events": survivor.detention_events.all(),
        "detention_periods": survivor.detention_periods.select_related("facility").all(),
        "witnesses": survivor.witnesses.select_related("facility_witnessed_at").all(),
        "documents": survivor.documents.all(),
        "medical_assessments": survivor.medical_assessments.all(),
        "interviews": survivor.interviews.all(),
        "today": datetime.now(),
    })


# ============================================================
# API لقوائم المدن والأحياء (لتغذية القوائم المتدرّجة)
# ============================================================

from django.http import JsonResponse
from .syria_geo import cities_for, neighborhoods_for


@login_required
def api_cities(request):
    gov = request.GET.get("gov", "")
    return JsonResponse({"cities": cities_for(gov)})


@login_required
def api_neighborhoods(request):
    gov = request.GET.get("gov", "")
    city = request.GET.get("city", "")
    return JsonResponse({"neighborhoods": neighborhoods_for(gov, city)})


# ============================================================
# تصدير/استيراد ملف ناجٍ مفرد (يشمل كل البيانات المرتبطة)
# ============================================================

@login_required
def export_single_survivor_json(request, pk):
    """تصدير ملف ناجٍ واحد بكل بياناته المرتبطة لملف JSON قابل للاستيراد."""
    from social_survey.models import (
        Child, EducationStatus, EmploymentInfo, HealthAccess,
        HousingInfo, HouseholdSurvey, NeedsAssessment,
    )
    survivor = get_object_or_404(SurvivorProfile.all_objects, pk=pk)

    # نجمع كل السجلات المرتبطة
    objects = [survivor]
    for related_attr in [
        "consent", "release_event", "long_term_impact",
    ]:
        obj = getattr(survivor, related_attr, None)
        if obj:
            objects.append(obj)
    for queryset_attr in [
        "detention_events", "detention_periods", "witnesses",
        "documents", "medical_assessments", "notes", "interviews",
    ]:
        objects.extend(list(getattr(survivor, queryset_attr).all()))
    # وسائط المقابلات + سلسلة الحيازة (متعدّية)
    for interview in survivor.interviews.all():
        objects.extend(list(interview.media.all()))
    for document in survivor.documents.all():
        objects.extend(list(document.custody_log.all()))
    # المسح الاجتماعي
    household = getattr(survivor, "household_survey", None)
    if household:
        objects.append(household)
        for related_attr in ["housing", "health_access", "needs"]:
            obj = getattr(household, related_attr, None)
            if obj:
                objects.append(obj)
        objects.extend(list(household.children.all()))
    education = getattr(survivor, "education", None)
    if education:
        objects.append(education)
    employment = getattr(survivor, "employment", None)
    if employment:
        objects.append(employment)

    data = serialize("json", objects, indent=2, ensure_ascii=False)
    AuditLog.objects.create(
        user=request.user, action=AuditLog.Action.EXPORT,
        target_model="SurvivorProfile", target_id=str(survivor.pk),
        target_repr=f"تصدير ملف مفرد {survivor.case_reference} ({len(objects)} سجل)",
        path=request.path, ip_address=request.META.get("REMOTE_ADDR"),
    )
    filename = f"survivor_{survivor.case_reference}_{datetime.now():%Y%m%d_%H%M%S}.json"
    response = HttpResponse(data, content_type="application/json; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
